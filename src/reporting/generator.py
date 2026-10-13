import json
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.exceptions import ReportGenerationError
from src.logger import get_logger
from src.processing.analytics import DataProcessor

EXCEL_PREFIX = "library_report"
CSV_PREFIX = "library_summary"
JSON_PREFIX = "library_export"

PRIMARY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=PRIMARY)
ZEBRA_FILL = PatternFill("solid", fgColor="EEF2F8")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=PRIMARY)
SUBTITLE_FONT = Font(italic=True, size=10, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MIN_WIDTH, MAX_WIDTH = 10, 60

KPI_LABELS = {
    "total_works": "Total works",
    "total_authors": "Total authors",
    "total_editions": "Total editions",
    "avg_editions_per_work": "Average editions per work",
    "median_publish_year": "Median first publish year",
    "oldest_publish_year": "Oldest first publish year",
    "newest_publish_year": "Newest first publish year",
    "isbn_coverage_pct": "ISBN coverage (%)",
    "distinct_languages": "Distinct languages",
    "top_language": "Top language",
    "top_subject": "Top subject",
    "top_author": "Top author",
}


def cell_value(value: Any) -> Any:
    """Converts pandas/numpy scalars into values openpyxl can serialize."""
    if value is None or (not isinstance(value, str | list | dict) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.tz_localize(None).to_pydatetime() if value.tzinfo else value.to_pydatetime()
    return value.item() if hasattr(value, "item") else value


def pretty_header(column: str) -> str:
    return column.replace("_", " ").replace("pct", "%").title().replace("Isbn", "ISBN")


@dataclass
class ReportBundle:
    excel: Path
    csv: Path
    json: Path
    kpis: dict[str, Any] = field(default_factory=dict)

    @property
    def files(self) -> list[Path]:
        return [self.excel, self.csv, self.json]


class ReportGenerator:
    def __init__(
        self,
        processor: DataProcessor,
        output_dir: Path,
        top_n: int = 10,
        now: datetime | None = None,
    ) -> None:
        self.processor = processor
        self.output_dir = Path(output_dir)
        self.top_n = top_n
        self.now = now or datetime.now()
        self.logger = get_logger("reporting")

    @property
    def stamp(self) -> str:
        return self.now.strftime("%Y%m%d")

    def path_for(self, prefix: str, extension: str) -> Path:
        return self.output_dir / f"{prefix}_{self.stamp}.{extension}"

    def _ensure_ready(self) -> None:
        if self.processor.is_empty:
            raise ReportGenerationError(
                "The database has no works yet. Ingest a subject or a catalog file first."
            )
        try:
            self.output_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise ReportGenerationError(f"Cannot create output folder: {exc}") from exc

    def datasets(self) -> dict[str, pd.DataFrame]:
        p, n = self.processor, self.top_n
        return {
            "Works": p.works_table(),
            "Top Authors": p.top_authors(n),
            "Publication Years": p.top_publication_years(n),
            "Languages": p.top_languages(n),
            "Subjects": p.top_subjects(n + 5),
            "Decades": p.works_by_decade(),
        }

    # ------------------------------------------------------------------ Excel
    def _write_kpi_sheet(self, sheet: Worksheet, kpis: dict[str, Any]) -> None:
        sheet.title = "KPIs"
        sheet["A1"] = "Open Library Intelligence - Key Performance Metrics"
        sheet["A1"].font = TITLE_FONT
        sheet["A2"] = f"Generated on {self.now:%Y-%m-%d %H:%M}"
        sheet["A2"].font = SUBTITLE_FONT
        sheet.merge_cells("A1:B1")

        start = 4
        for column, header in enumerate(("Metric", "Value"), start=1):
            cell = sheet.cell(row=start, column=column, value=header)
            cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER

        for offset, (key, label) in enumerate(KPI_LABELS.items(), start=1):
            row = start + offset
            value = kpis.get(key)
            sheet.cell(row=row, column=1, value=label).border = BORDER
            value_cell = sheet.cell(row=row, column=2, value="N/A" if value is None else value)
            value_cell.border = BORDER
            value_cell.alignment = Alignment(horizontal="right")
            if isinstance(value, float):
                value_cell.number_format = "#,##0.00"
            elif isinstance(value, int) and not key.endswith("_year"):
                value_cell.number_format = "#,##0"
            if offset % 2 == 0:
                for column in (1, 2):
                    sheet.cell(row=row, column=column).fill = ZEBRA_FILL
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 28

    def _write_table_sheet(self, sheet: Worksheet, frame: pd.DataFrame) -> None:
        columns = list(frame.columns)
        for index, column in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index, value=pretty_header(str(column)))
            cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_index, row in enumerate(frame.itertuples(index=False), start=2):
            zebra = row_index % 2 == 1
            for col_index, value in enumerate(row, start=1):
                cell = sheet.cell(row=row_index, column=col_index, value=cell_value(value))
                cell.border = BORDER
                if zebra:
                    cell.fill = ZEBRA_FILL
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"

        for index, column in enumerate(columns, start=1):
            values = [pretty_header(str(column)), *(str(v) for v in frame[column].head(500))]
            width = max(len(value) for value in values) + 2
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = max(MIN_WIDTH, min(width, MAX_WIDTH))

        sheet.freeze_panes = "A2"
        if columns:
            last = get_column_letter(len(columns))
            sheet.auto_filter.ref = f"A1:{last}{max(len(frame) + 1, 2)}"

    def generate_excel(self, kpis: dict[str, Any] | None = None) -> Path:
        self._ensure_ready()
        kpis = kpis if kpis is not None else self.processor.kpis()
        path = self.path_for(EXCEL_PREFIX, "xlsx")
        workbook = Workbook()
        self._write_kpi_sheet(workbook.active, kpis)
        for title, frame in self.datasets().items():
            self._write_table_sheet(workbook.create_sheet(title), frame)
        try:
            workbook.save(path)
        except PermissionError as exc:
            raise ReportGenerationError(
                f"Cannot write {path.name}. Close it if it is open in Excel and try again."
            ) from exc
        except OSError as exc:
            raise ReportGenerationError(f"Failed to save Excel report: {exc}") from exc
        self.logger.info("Excel report written to %s", path)
        return path

    # -------------------------------------------------------------------- CSV
    def generate_csv(self) -> Path:
        self._ensure_ready()
        path = self.path_for(CSV_PREFIX, "csv")
        try:
            self.processor.summary(self.top_n).to_csv(path, index=False, encoding="utf-8")
        except OSError as exc:
            raise ReportGenerationError(f"Failed to save CSV summary: {exc}") from exc
        self.logger.info("CSV summary written to %s", path)
        return path

    # ------------------------------------------------------------------- JSON
    def build_payload(self, kpis: dict[str, Any] | None = None) -> dict[str, Any]:
        p, to_records = self.processor, DataProcessor.to_records
        return {
            "generated_at": self.now.isoformat(timespec="seconds"),
            "kpis": kpis if kpis is not None else p.kpis(),
            "top_publication_years": to_records(p.top_publication_years(self.top_n)),
            "top_languages": to_records(p.top_languages(self.top_n)),
            "top_subjects": to_records(p.top_subjects(self.top_n)),
            "top_authors": to_records(p.top_authors(self.top_n)),
            "works": to_records(p.works_table()),
        }

    def generate_json(self, kpis: dict[str, Any] | None = None) -> Path:
        self._ensure_ready()
        path = self.path_for(JSON_PREFIX, "json")
        payload = self.build_payload(kpis)
        try:
            path.write_text(
                json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
            )
        except OSError as exc:
            raise ReportGenerationError(f"Failed to save JSON export: {exc}") from exc
        self.logger.info("JSON export written to %s", path)
        return path

    def generate_all(self) -> ReportBundle:
        self._ensure_ready()
        kpis = self.processor.kpis()
        return ReportBundle(
            excel=self.generate_excel(kpis),
            csv=self.generate_csv(),
            json=self.generate_json(kpis),
            kpis=kpis,
        )
